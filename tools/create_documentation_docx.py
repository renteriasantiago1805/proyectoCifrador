from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Documentacion_proyecto_cifrado.docx"


def style_run(run, size=10.5, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        style_run(run, 13 if level == 1 else 11.5, True)
    return paragraph


def add_paragraph(document, text="", bold=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    style_run(run, 10.5, bold)
    return paragraph


def add_list_item(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    style_run(run, 10.5)
    return paragraph


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, 9.2, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_reference(document, number, function_name, how_it_works, purpose):
    add_heading(document, f"Referencia {number}: {function_name}", 2)
    add_paragraph(document, "Como trabaja:", True)
    add_paragraph(document, how_it_works)
    add_paragraph(document, "Para que sirve:", True)
    add_paragraph(document, purpose)


def add_prompt_table(document):
    add_heading(document, "Prompts representativos y recursos utilizados", 1)
    add_paragraph(
        document,
        "Durante el desarrollo se usaron instrucciones iterativas para construir, corregir, publicar y documentar el sistema. No se incluyen todos los mensajes textuales, sino los prompts mas significativos agrupados por la intencion de trabajo.",
    )

    prompt_rows = [
        (
            "Prompt 1",
            "Desarrollar una aplicacion en Angular y Bootstrap para cifrar y descifrar mensajes con Cesar y Atbash, siguiendo la rubrica del proyecto.",
            "Definio la base del proyecto: tecnologias, alcance, metodos de cifrado, descifrado automatico, publicacion web y documentacion segura.",
        ),
        (
            "Prompt 2",
            "Ajustar la interfaz para que fuera mas compacta, sin titulo innecesario, con el area de cifrado y descifrado visibles y el conjunto de caracteres en una zona inferior.",
            "Guio el diseno visual final para que la aplicacion fuera directa, facil de presentar y sin elementos que pertenecieran mas al documento que a la pagina.",
        ),
        (
            "Prompt 3",
            "Corregir el cifrado Cesar para que funcionara con modulos negativos y con rangos validos segun el tamano del alfabeto.",
            "Permitio mejorar la logica modular y mostrar desplazamientos equivalentes de forma mas clara.",
        ),
        (
            "Prompt 4",
            "Permitir caracteres raros en el cifrado cuando estuvieran dentro del conjunto de caracteres, pero ignorarlos durante el descifrado automatico si no pertenecian al alfabeto activo.",
            "Definio el manejo de alfabetos personalizados y evito que simbolos externos afectaran la deteccion automatica.",
        ),
        (
            "Prompt 5",
            "Documentar el codigo de forma segura usando comentarios de referencia y explicar cada referencia en un documento Word.",
            "Produjo el esquema de Referencia 1, Referencia 2, etc., separando el codigo de la explicacion detallada para no llenar los archivos fuente de comentarios extensos.",
        ),
        (
            "Prompt 6",
            "Publicar el programa en GitHub Pages, verificar que funcionara en la liga publica y corregir diferencias entre la version local y la version publicada.",
            "Aseguro que la entrega tuviera enlace funcional y que GitHub Pages cargara el mismo build de Angular que se veia en local.",
        ),
        (
            "Prompt 7",
            "Usar como referencia una documentacion en PDF de otro proyecto, sin copiarla, para que el reporte tuviera una estructura mas formal.",
            "Ayudo a reforzar la organizacion del documento con portada, indice, fundamento tecnico, pruebas, enlaces y bibliografia, conservando redaccion propia.",
        ),
    ]

    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = table.rows[0].cells
    set_cell_text(headers[0], "Prompt", True)
    set_cell_text(headers[1], "Instruccion representativa", True)
    set_cell_text(headers[2], "Uso dentro del proyecto", True)
    for label, prompt, usage in prompt_rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, True)
        set_cell_text(cells[1], prompt)
        set_cell_text(cells[2], usage)

    add_heading(document, "Recursos utilizados", 2)
    for item in [
        "Rubrica e instrucciones del proyecto: se usaron como guia principal para cubrir portada, indice, introduccion, objetivo, desarrollo, publicacion, conclusion y bibliografia.",
        "Angular: se uso para construir la aplicacion web, separar componentes y mantener la logica de cifrado en archivos TypeScript.",
        "Bootstrap: se uso para formularios, botones, rejilla responsiva y estilos base de la interfaz.",
        "GitHub y GitHub Pages: se usaron para alojar el codigo fuente y publicar la version web del programa.",
        "Documento PDF de referencia: se reviso como ejemplo de estructura de reporte tecnico, sin copiar texto ni codigo.",
        "Pruebas locales: se usaron npm test y npm run build para validar que el proyecto compilara y que las pruebas del cifrado pasaran.",
        "Navegador local y GitHub Pages: se usaron para verificar que la aplicacion cifrara, descifrara y mostrara correctamente el resultado publicado.",
    ]:
        add_list_item(document, item)


references = [
    (
        1,
        "defaultCharacters",
        "Construye el alfabeto inicial con los caracteres ASCII imprimibles. Para hacerlo recorre los codigos del 32 al 126 y convierte cada numero a su caracter correspondiente. El resultado se une en una sola cadena que se puede mostrar y editar en la interfaz.",
        "Sirve como base del sistema cuando el usuario no define un conjunto propio. Con este alfabeto el programa puede cifrar letras, numeros, espacios y signos comunes sin depender de valores secretos.",
    ),
    (
        2,
        "encryptCaesar",
        "Recibe el texto, el conjunto de caracteres y el modulo. Primero limpia caracteres repetidos del alfabeto, despues normaliza el modulo con safeShift y finalmente manda cada caracter a transformByIndex para moverlo dentro del alfabeto.",
        "Sirve para cifrar con Cesar usando modulos positivos, negativos o mayores al tamano del alfabeto. Asi el usuario puede probar desplazamientos distintos sin que el programa falle.",
    ),
    (
        3,
        "encryptAtbash y decryptAtbash",
        "Atbash toma la posicion de cada caracter dentro del alfabeto y la cambia por la posicion opuesta. Por ejemplo, el primer caracter se cambia por el ultimo, el segundo por el penultimo y asi sucesivamente. Como Atbash es reversible, la misma logica cifra y descifra.",
        "Sirve para aplicar el segundo metodo solicitado en la rubrica. Tambien permite comparar un cifrado por desplazamiento contra un cifrado por sustitucion reflejada.",
    ),
    (
        4,
        "detectAndDecrypt",
        "Genera una lista de posibles respuestas: una usando Atbash y una por cada modulo posible de Cesar. Antes de probar candidatos elimina del texto cifrado los caracteres que no pertenecen al alfabeto elegido, porque esos simbolos no tienen posicion valida. Luego calcula una puntuacion para cada candidato y conserva solo el mejor.",
        "Sirve para descifrar sin que el usuario elija manualmente si el mensaje esta en Cesar o Atbash. Si el mejor resultado es Cesar, tambien devuelve el modulo detectado.",
    ),
    (
        5,
        "transformByIndex",
        "Convierte el texto en caracteres individuales, busca cada uno dentro del alfabeto y mueve su posicion segun el modulo recibido. Si el caracter no esta en el alfabeto, lo deja igual en cifrado y descifrado directo.",
        "Sirve como funcion central para el movimiento modular de Cesar. Tambien evita perdida de informacion cuando el usuario escribe simbolos que no estan incluidos en el conjunto de caracteres.",
    ),
    (
        6,
        "scoreSpanishPlainText",
        "Normaliza el candidato a minusculas, retira acentos para comparar mejor y analiza si el resultado se parece al espanol. La puntuacion sube cuando aparecen palabras comunes, vocales en proporcion normal, espacios reales y estructura de frase legible. La puntuacion baja si hay secuencias poco naturales.",
        "Sirve para incorporar el conocimiento de Al-Kindi. En lugar de pedir al usuario que revise todas las posibilidades, el programa usa patrones del idioma para elegir automaticamente la linea descifrada mas probable.",
    ),
    (
        7,
        "scoreReadableSeparators",
        "Revisa si el texto candidato conserva espacios normales entre palabras y penaliza separadores artificiales como guiones bajos, diagonales invertidas o simbolos que suelen aparecer cuando una prueba no produjo una frase natural.",
        "Sirve para mejorar la deteccion automatica y reducir falsos positivos, especialmente cuando hay alfabetos grandes o con muchos simbolos.",
    ),
    (
        8,
        "toSignedShift",
        "Convierte un modulo valido a una forma mas facil de leer. Si el desplazamiento detectado es mayor que la mitad del alfabeto, lo expresa como equivalente negativo. Por ejemplo, un modulo grande puede representarse como -3 cuando ambos producen el mismo resultado.",
        "Sirve para que la interfaz muestre el modulo de Cesar de forma comprensible y tambien para soportar correctamente modulos negativos.",
    ),
    (
        9,
        "missingCharacters",
        "Compara los caracteres de la oracion original contra el alfabeto activo. Los que no existan en el conjunto se guardan una sola vez para avisar al usuario y permitir agregarlos al alfabeto.",
        "Sirve para que el cifrado pueda trabajar con letras acentuadas, signos o simbolos raros cuando el usuario los quiera incluir. Esta validacion se enfoca en el texto a cifrar, no en el texto cifrado automatico.",
    ),
    (
        10,
        "rarePreset y useRarePreset",
        "rarePreset guarda un conjunto largo de simbolos no comunes. useRarePreset copia ese conjunto al campo de caracteres y ajusta el rango permitido del modulo para que coincida con el tamano del alfabeto.",
        "Sirve para probar que el programa puede cifrar usando caracteres especiales cuando el usuario decide incluirlos en el conjunto de trabajo.",
    ),
    (
        11,
        "textUsedForDecryption y charactersOutsideAlphabet",
        "textUsedForDecryption crea una version del texto cifrado que solo conserva caracteres del alfabeto seleccionado y espacios. charactersOutsideAlphabet registra que caracteres fueron ignorados para poder mostrarlos en pantalla.",
        "Sirve para cumplir la regla acordada del descifrado: si aparecen caracteres raros que no estan en el alfabeto elegido, el sistema los elimina e intenta descifrar con el alfabeto activo.",
    ),
]


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Reporte tecnico del cifrador Cesar y Atbash")
style_run(run, 16, True)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Aplicacion web en Angular Bootstrap y analisis automatico de frecuencia")
style_run(run, 11)

add_heading(document, "Portada", 1)
cover_rows = [
    ("Materia", "Seguridad informatica"),
    ("Proyecto", "Programa web para cifrar y descifrar oraciones con Cesar y Atbash"),
    ("Integrantes", "[Agregar nombres]"),
    ("Profesor", "[Agregar nombre del profesor]"),
    ("Grupo", "[Agregar grupo]"),
    ("Fecha", "[Agregar fecha de entrega]"),
    ("Programa web", "https://renteriasantiago1805.github.io/proyectoCifrador/"),
    ("Codigo fuente", "https://github.com/renteriasantiago1805/proyectoCifrador"),
]
cover_table = document.add_table(rows=1, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_table.style = "Table Grid"
set_cell_text(cover_table.rows[0].cells[0], "Campo", True)
set_cell_text(cover_table.rows[0].cells[1], "Informacion", True)
for key, value in cover_rows:
    row = cover_table.add_row().cells
    set_cell_text(row[0], key, True)
    set_cell_text(row[1], value)

add_heading(document, "Indice general", 1)
for item in [
    "1. Introduccion",
    "2. Objetivo",
    "3. Fundamento tecnico",
    "4. Desarrollo del programa",
    "5. Indice de referencias del codigo",
    "6. Explicacion detallada de referencias",
    "7. Pruebas de funcionamiento",
    "8. Enlaces de entrega",
    "9. Prompts representativos y recursos utilizados",
    "10. Conclusion",
    "11. Bibliografia",
]:
    add_paragraph(document, item)

add_heading(document, "Introduccion", 1)
add_paragraph(
    document,
    "Abu Yusuf Yaqub ibn Ishaq al-Kindi, أبو يوسف يعقوب بن إسحاق الكندي, aporto una de las bases historicas del criptoanalisis: el analisis de frecuencia. Su idea principal consiste en observar que cada idioma repite ciertas letras, combinaciones y palabras con mayor frecuencia. Al comparar esos patrones con un texto cifrado, es posible proponer descifrados probables sin conocer la clave original.",
)
add_paragraph(
    document,
    "Los cifrados Cesar y Atbash son utiles para aprender sustitucion, modulo y alfabetos de trabajo, pero no son viables para proteger datos reales. Ambos mantienen patrones del idioma, tienen un espacio de busqueda pequeno y pueden romperse con pruebas automatizadas. Por eso en sistemas modernos se usan algoritmos criptograficos con claves robustas, analisis publico, administracion de llaves y resistencia matematica mucho mayor.",
)

add_heading(document, "Objetivo", 1)
add_paragraph(
    document,
    "Desarrollar una aplicacion web en Angular y Bootstrap que permita cifrar mensajes con Cesar o Atbash, usar un conjunto de caracteres configurable y descifrar automaticamente identificando el tipo de cifrado y el modulo cuando corresponde.",
)

add_heading(document, "Fundamento tecnico", 1)
add_paragraph(
    document,
    "El programa trabaja con el conjunto de caracteres como si fuera un alfabeto ordenado. Cada simbolo tiene una posicion, y esa posicion es la que se usa para aplicar las operaciones matematicas. Esto permite usar ASCII, letras en espanol o simbolos especiales sin depender de que sus codigos Unicode sean consecutivos.",
)
add_paragraph(document, "Cifrado Cesar:", True)
add_paragraph(
    document,
    "Si C es el conjunto de caracteres, n su longitud, idx(x) la posicion del caracter x y k el modulo, el cifrado se calcula como C[(idx(x) + k) mod n]. Para descifrar directamente se usa C[(idx(x) - k) mod n]. El modulo evita que el indice se salga del alfabeto y permite regresar al inicio cuando se llega al final.",
)
add_paragraph(document, "Cifrado Atbash:", True)
add_paragraph(
    document,
    "Atbash no usa una clave numerica. Su regla consiste en reflejar cada posicion: C[(n - 1) - idx(x)]. Por esa razon el mismo procedimiento sirve para cifrar y descifrar, ya que aplicar la reflexion dos veces devuelve el texto original.",
)
add_paragraph(document, "Analisis automatico:", True)
add_paragraph(
    document,
    "Para el descifrado automatico se prueban todos los candidatos posibles. Despues se puntuan con caracteristicas del espanol: palabras frecuentes, proporcion de vocales, espacios reales y penalizacion de secuencias poco naturales. Esta parte representa la aplicacion practica del analisis de frecuencia asociado con Al-Kindi.",
)

add_heading(document, "Desarrollo del programa", 1)
add_paragraph(
    document,
    "La aplicacion esta construida con Angular para separar la logica del cifrado de la interfaz. Bootstrap se utiliza para organizar los controles visuales, formularios, botones y resultados. La logica principal vive en src/app/crypto.ts, mientras que la interaccion del usuario vive en src/app/cifrador/cifrador.ts y su plantilla HTML.",
)
add_paragraph(
    document,
    "El usuario puede escribir el texto original, seleccionar Cesar o Atbash, indicar el modulo en Cesar y elegir el alfabeto de trabajo. El alfabeto puede ser ASCII, espanol ampliado o un conjunto con simbolos especiales. El sistema elimina caracteres repetidos para que cada simbolo tenga una unica posicion dentro del modulo.",
)
add_paragraph(
    document,
    "Para cifrar con Cesar, el programa desplaza cada caracter dentro del alfabeto mediante aritmetica modular. Para Atbash, sustituye cada caracter por su equivalente simetrico. Para descifrar automaticamente, genera posibles soluciones y las califica con reglas inspiradas en el analisis de frecuencia de Al-Kindi.",
)
add_paragraph(
    document,
    "En descifrado automatico, si el texto contiene caracteres que no estan en el alfabeto seleccionado, esos caracteres se ignoran antes de evaluar los candidatos. Esto evita que simbolos externos alteren la deteccion y permite continuar con el alfabeto activo.",
)

add_heading(document, "Indice de referencias del codigo", 1)
add_paragraph(
    document,
    "En el codigo solo aparecen comentarios con el formato Referencia N. La explicacion completa se concentra en esta documentacion para mantener el codigo limpio y documentado de forma segura.",
)
reference_table = document.add_table(rows=1, cols=3)
reference_table.alignment = WD_TABLE_ALIGNMENT.CENTER
reference_table.style = "Table Grid"
headers = reference_table.rows[0].cells
set_cell_text(headers[0], "Referencia", True)
set_cell_text(headers[1], "Funcion o bloque", True)
set_cell_text(headers[2], "Archivo", True)

file_by_reference = {
    1: "src/app/crypto.ts",
    2: "src/app/crypto.ts",
    3: "src/app/crypto.ts",
    4: "src/app/crypto.ts",
    5: "src/app/crypto.ts",
    6: "src/app/crypto.ts",
    7: "src/app/crypto.ts",
    8: "src/app/crypto.ts",
    9: "src/app/cifrador/cifrador.ts",
    10: "src/app/cifrador/cifrador.ts",
    11: "src/app/crypto.ts",
}
for number, function_name, _how, _purpose in references:
    cells = reference_table.add_row().cells
    set_cell_text(cells[0], f"Referencia {number}", True)
    set_cell_text(cells[1], function_name)
    set_cell_text(cells[2], file_by_reference[number])

add_heading(document, "Explicacion detallada de referencias", 1)
for reference in references:
    add_reference(document, *reference)

add_heading(document, "Pruebas de funcionamiento", 1)
add_list_item(document, "Prueba 1: escribir una oracion, elegir Cesar y usar modulo 3. El resultado debe cambiar y el descifrado automatico debe regresar la frase original.")
add_list_item(document, "Prueba 2: escribir una oracion, elegir Cesar y usar modulo -3. El descifrado automatico debe indicar Cesar con modulo -3 o su equivalente modular.")
add_list_item(document, "Prueba 3: elegir Atbash. Como es reversible, el descifrado automatico debe reconocer Atbash y mostrar la oracion original.")
add_list_item(document, "Prueba 4: escribir un caracter que no este en el alfabeto al cifrar. El sistema debe avisar que falta y permitir agregarlo.")
add_list_item(document, "Prueba 5: pegar un texto cifrado con simbolos externos al alfabeto en descifrado automatico. El sistema debe ignorar esos simbolos, informar cuales elimino y continuar con la deteccion.")
add_list_item(document, "Prueba 6: usar el conjunto raro de caracteres para cifrar. El resultado debe incluir simbolos especiales y el modulo debe mantenerse dentro del rango valido.")

add_heading(document, "Enlaces de entrega", 1)
add_paragraph(document, "Programa web publicado: https://renteriasantiago1805.github.io/proyectoCifrador/")
add_paragraph(document, "Codigo fuente documentado: https://github.com/renteriasantiago1805/proyectoCifrador")

add_prompt_table(document)

add_heading(document, "Conclusion", 1)
add_paragraph(
    document,
    "El proyecto demuestra el funcionamiento de dos cifrados clasicos y tambien sus debilidades. Al automatizar el analisis de frecuencia, se observa que Cesar y Atbash pueden romperse sin intervencion humana directa, por lo que no son adecuados para proteger informacion real. La aplicacion cumple como herramienta educativa porque permite experimentar con alfabetos, modulos y deteccion automatica de una manera visual.",
)

add_heading(document, "Bibliografia", 1)
for source in [
    "1001 Inventions. (s. f.). Code breaking a thousand years ago. https://www.1001inventions.com/code-breaking/",
    "Khan Academy. (s. f.). Ancient cryptography. https://www.khanacademy.org/computing/computer-science/cryptography",
    "Wikipedia contributors. (s. f.). Atbash. Wikipedia. https://en.wikipedia.org/wiki/Atbash",
    "Wikipedia contributors. (s. f.). Caesar cipher. Wikipedia. https://en.wikipedia.org/wiki/Caesar_cipher",
    "Wikipedia contributors. (s. f.). Frequency analysis. Wikipedia. https://en.wikipedia.org/wiki/Frequency_analysis",
]:
    add_paragraph(document, source)

document.save(OUTPUT)
print(OUTPUT)
